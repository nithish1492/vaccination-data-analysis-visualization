"""Rebuild vaccination.db from the supplied Excel files.
Run from the project root: python scripts/clean_data.py
Missing values are preserved as NULL; metadata rows without YEAR are excluded.
"""
import os, sqlite3, math
from openpyxl import load_workbook

ROOT=os.path.dirname(os.path.dirname(os.path.abspath(__file__))); DB=os.path.join(ROOT,"vaccination.db")
S=[
("coverage","coverage-data.xlsx",['GROUP','CODE','NAME','YEAR','ANTIGEN','ANTIGEN_DESCRIPTION','COVERAGE_CATEGORY','COVERAGE_CATEGORY_DESCRIPTION','TARGET_NUMBER','DOSES','COVERAGE'],['group_name','code','name','year','antigen','antigen_description','coverage_category','coverage_category_description','target_number','doses','coverage'],('CODE','NAME','GROUP')),
("incidence_rate","incidence-rate-data(1).xlsx",['GROUP','CODE','NAME','YEAR','DISEASE','DISEASE_DESCRIPTION','DENOMINATOR','INCIDENCE_RATE'],['group_name','code','name','year','disease','disease_description','denominator','incidence_rate'],('CODE','NAME','GROUP')),
("reported_cases","reported-cases-data(1).xlsx",['GROUP','CODE','NAME','YEAR','DISEASE','DISEASE_DESCRIPTION','CASES'],['group_name','code','name','year','disease','disease_description','cases'],('CODE','NAME','GROUP')),
("vaccine_introduction","vaccine-introduction-data(1).xlsx",['ISO_3_CODE','COUNTRYNAME','WHO_REGION','YEAR','DESCRIPTION','INTRO'],['iso_3_code','country_name','who_region','year','description','intro'],('ISO_3_CODE','COUNTRYNAME',None)),
("vaccine_schedule","vaccine-schedule-data(1).xlsx",['ISO_3_CODE','COUNTRYNAME','WHO_REGION','YEAR','VACCINECODE','VACCINE_DESCRIPTION','SCHEDULEROUNDS','TARGETPOP','TARGETPOP_DESCRIPTION','GEOAREA','AGEADMINISTERED','SOURCECOMMENT'],['iso_3_code','country_name','who_region','year','vaccine_code','vaccine_description','schedule_rounds','target_pop','target_pop_description','geoarea','age_administered','source_comment'],('ISO_3_CODE','COUNTRYNAME',None))]
def norm(v):
    if isinstance(v,float) and math.isnan(v): return None
    if isinstance(v,str): return v.strip() or None
    return v
def main():
    if os.path.exists(DB): os.remove(DB)
    db=sqlite3.connect(DB); c=db.cursor()
    c.executescript("""PRAGMA journal_mode=OFF;PRAGMA synchronous=OFF;
    CREATE TABLE geography(code TEXT PRIMARY KEY,name TEXT NOT NULL,group_name TEXT);
    CREATE TABLE countries(code TEXT PRIMARY KEY,name TEXT NOT NULL,group_name TEXT);
    CREATE TABLE coverage(id INTEGER PRIMARY KEY AUTOINCREMENT,group_name TEXT,code TEXT,name TEXT,year INTEGER,antigen TEXT,antigen_description TEXT,coverage_category TEXT,coverage_category_description TEXT,target_number REAL,doses REAL,coverage REAL);
    CREATE TABLE incidence_rate(id INTEGER PRIMARY KEY AUTOINCREMENT,group_name TEXT,code TEXT,name TEXT,year INTEGER,disease TEXT,disease_description TEXT,denominator TEXT,incidence_rate REAL);
    CREATE TABLE reported_cases(id INTEGER PRIMARY KEY AUTOINCREMENT,group_name TEXT,code TEXT,name TEXT,year INTEGER,disease TEXT,disease_description TEXT,cases REAL);
    CREATE TABLE vaccine_introduction(id INTEGER PRIMARY KEY AUTOINCREMENT,iso_3_code TEXT,country_name TEXT,who_region TEXT,year INTEGER,description TEXT,intro TEXT);
    CREATE TABLE vaccine_schedule(id INTEGER PRIMARY KEY AUTOINCREMENT,iso_3_code TEXT,country_name TEXT,who_region TEXT,year INTEGER,vaccine_code TEXT,vaccine_description TEXT,schedule_rounds INTEGER,target_pop TEXT,target_pop_description TEXT,geoarea TEXT,age_administered TEXT,source_comment TEXT);
    CREATE TABLE data_quality_log(table_name TEXT PRIMARY KEY,source_file TEXT,source_rows INTEGER,cleaned_rows INTEGER,duplicates_removed INTEGER,missing_key_rows INTEGER,notes TEXT);""")
    geo={}
    for t,file,src,dst,g in S:
        path=os.path.join(ROOT,"data",file)
        if not os.path.exists(path): path=os.path.join(ROOT,file)
        wb=load_workbook(path,read_only=True,data_only=True); ws=wb["Data"]; it=ws.iter_rows(values_only=True); hdr=next(it); idx={h:i for i,h in enumerate(hdr)}
        for r in it:
            vals=[norm(r[idx[h]]) if idx[h]<len(r) else None for h in src]
            if vals[3] is None or vals[0] is None or vals[1] is None: continue
            grp=vals[src.index(g[2])] if g[2] else "COUNTRIES"
            geo.setdefault(vals[1],(vals[2],grp))
        wb.close()
    c.executemany("INSERT INTO geography VALUES(?,?,?)",[(k,v[0],v[1]) for k,v in geo.items()])
    c.execute("INSERT OR IGNORE INTO countries SELECT code,name,group_name FROM geography WHERE length(code)=3 AND group_name='COUNTRIES'")
    db.commit()
    for t,file,src,dst,g in S:
        path=os.path.join(ROOT,"data",file)
        if not os.path.exists(path): path=os.path.join(ROOT,file)
        wb=load_workbook(path,read_only=True,data_only=True); ws=wb["Data"]; it=ws.iter_rows(values_only=True); hdr=next(it); idx={h:i for i,h in enumerate(hdr)}
        sql=f"INSERT INTO {t} ({','.join(dst)}) VALUES ({','.join('?' for _ in dst)})"; buf=[]; total=cleaned=drop=0
        for r in it:
            total+=1; vals=[norm(r[idx[h]]) if idx[h]<len(r) else None for h in src]
            if vals[3] is None: drop+=1; continue
            buf.append(vals); cleaned+=1
            if len(buf)>=10000: c.executemany(sql,buf); buf=[]
        if buf:c.executemany(sql,buf)
        c.execute("INSERT INTO data_quality_log VALUES(?,?,?,?,?,?,?)",(t,file,total,cleaned,0,drop,"Trimmed text; blank values -> NULL; metadata rows without YEAR excluded."))
        db.commit(); wb.close()
    c.executescript("""CREATE INDEX idx_cov_code_year ON coverage(code,year);CREATE INDEX idx_cov_antigen_year ON coverage(antigen,year);CREATE INDEX idx_cov_category ON coverage(coverage_category);
    CREATE INDEX idx_inc_code_year ON incidence_rate(code,year);CREATE INDEX idx_inc_disease_year ON incidence_rate(disease,year);
    CREATE INDEX idx_cases_code_year ON reported_cases(code,year);CREATE INDEX idx_cases_disease_year ON reported_cases(disease,year);
    CREATE INDEX idx_intro_country_year ON vaccine_introduction(iso_3_code,year);CREATE INDEX idx_intro_desc_year ON vaccine_introduction(description,year);CREATE INDEX idx_sched_country_year ON vaccine_schedule(iso_3_code,year);""")
    # Trigger-based referential integrity because source contains country, regional and global codes.
    c.executescript("""CREATE TRIGGER coverage_geo_fk_insert BEFORE INSERT ON coverage WHEN NEW.code IS NOT NULL AND NOT EXISTS(SELECT 1 FROM geography WHERE code=NEW.code) BEGIN SELECT RAISE(ABORT,'Invalid geography code'); END;
    CREATE TRIGGER incidence_geo_fk_insert BEFORE INSERT ON incidence_rate WHEN NEW.code IS NOT NULL AND NOT EXISTS(SELECT 1 FROM geography WHERE code=NEW.code) BEGIN SELECT RAISE(ABORT,'Invalid geography code'); END;
    CREATE TRIGGER cases_geo_fk_insert BEFORE INSERT ON reported_cases WHEN NEW.code IS NOT NULL AND NOT EXISTS(SELECT 1 FROM geography WHERE code=NEW.code) BEGIN SELECT RAISE(ABORT,'Invalid geography code'); END;
    CREATE TRIGGER intro_geo_fk_insert BEFORE INSERT ON vaccine_introduction WHEN NEW.iso_3_code IS NOT NULL AND NOT EXISTS(SELECT 1 FROM geography WHERE code=NEW.iso_3_code) BEGIN SELECT RAISE(ABORT,'Invalid geography code'); END;
    CREATE TRIGGER schedule_geo_fk_insert BEFORE INSERT ON vaccine_schedule WHEN NEW.iso_3_code IS NOT NULL AND NOT EXISTS(SELECT 1 FROM geography WHERE code=NEW.iso_3_code) BEGIN SELECT RAISE(ABORT,'Invalid geography code'); END;""")
    db.commit(); print("Created",DB); print("Rows:",[(t,c.execute(f"select count(*) from {t}").fetchone()[0]) for t in ['countries','coverage','incidence_rate','reported_cases','vaccine_introduction','vaccine_schedule']]); db.close()
if __name__=="__main__": main()
